from dataclasses import replace
from io import BytesIO

from openpyxl import load_workbook

from company_gtm_enricher.people.excel_export import excel_bytes
from company_gtm_enricher.people.models import Candidate, CompanyTarget, DiscoveryBatch, PeopleCriteria
from company_gtm_enricher.people.providers import DiscoveryFailure
from company_gtm_enricher.people.service import PeopleService


def supported_candidate(company="Example"):
    return Candidate(
        name="Alex Example", title="Co-founder", is_founder=True,
        linkedin_url="https://www.linkedin.com/in/alex-example",
        evidence_url="https://www.linkedin.com/in/alex-example",
        evidence_text=f"Alex Example - Co-founder at {company} - LinkedIn",
        observed_at="2026-09-16T12:00:00+00:00",
        identity_verified=True, employment_verified=True, role_verified=True,
    )


class FailingDiscoveryProvider:
    is_mock = False

    def __init__(self, *, fatal=False, candidates=None):
        self.fatal = fatal
        self.candidates = candidates if candidates is not None else [
            replace(supported_candidate("Failure"), name="", title="", is_founder=False,
                    identity_verified=False, employment_verified=False, role_verified=False,
                    review_reason="Role classification failed."),
        ]
        self.calls = []
        # Prior usage belongs to earlier runs and must not leak into this result.
        self.usage = {"search_requests": 7, "llm_requests": 3}

    def discover(self, company, criteria):
        self.calls.append(company.name)
        self.usage["search_requests"] += 2
        self.usage["llm_requests"] += 1
        if company.name == "Failure":
            raise DiscoveryFailure(
                "Role classification timed out.",
                DiscoveryBatch(self.candidates, notes="Two searches completed before classification failed.", limited=True),
                fatal=self.fatal,
            )
        return [supported_candidate(company.name)]


def test_nonfatal_failure_retains_source_evidence_and_continues_to_next_company():
    provider = FailingDiscoveryProvider()
    snapshots = []
    result = PeopleService(provider).run(
        [CompanyTarget("Failure", "failure.example"), CompanyTarget("Next", "next.example")],
        PeopleCriteria(), on_update=lambda *update: snapshots.append(update),
    )

    assert result.status == "completed_with_errors"
    assert provider.calls == ["Failure", "Next"]
    assert result.provider_calls == 2
    assert result.usage == {"search_requests": 4, "llm_requests": 2}
    assert [row["Status"] for row in result.companies] == ["error", "processed"]
    assert result.companies[0]["People"] == 0
    assert result.companies[0]["Review"] == 1
    assert "timed out" in result.companies[0]["Notes"]
    assert "Two searches completed" in result.companies[0]["Notes"]
    assert len(result.people) == len(result.review) == 1
    review = result.review[0]
    assert review["Company"] == "Failure"
    assert review["Evidence Text"] == provider.candidates[0].evidence_text
    assert review["Evidence URL"] == provider.candidates[0].evidence_url
    assert review["Observed At"] == provider.candidates[0].observed_at
    assert review["Status"] == "processing_failed"
    assert "source evidence is retained for retry" in review["Notes"]
    assert "Role classification failed." in review["Review Reason"]
    assert "Name or title is missing." in review["Review Reason"]
    assert snapshots[-1][1:3] == (2, 2)
    assert snapshots[-1][0].status == "completed_with_errors"
    assert not snapshots[0][0].review
    result.review[0]["Evidence Text"] = "Changed after callback"
    assert snapshots[-1][0].review[0]["Evidence Text"] == provider.candidates[0].evidence_text


def test_fatal_failure_preserves_prior_people_current_review_and_future_pending():
    provider = FailingDiscoveryProvider(fatal=True)
    snapshots = []
    result = PeopleService(provider).run(
        [CompanyTarget("First"), CompanyTarget("Failure"), CompanyTarget("Later")],
        PeopleCriteria(), on_update=lambda *update: snapshots.append(update),
    )

    assert result.status == "failed"
    assert provider.calls == ["First", "Failure"]
    assert result.provider_calls == 2
    assert result.usage == {"search_requests": 4, "llm_requests": 2}
    assert [row["Status"] for row in result.companies] == ["processed", "error", "pending"]
    assert [row["Company"] for row in result.people] == ["First"]
    assert [row["Company"] for row in result.review] == ["Failure"]
    assert result.companies[1]["Review"] == 1
    assert result.companies[2]["People"] == result.companies[2]["Review"] == 0
    assert snapshots[-1][0].status == "failed"
    assert snapshots[-1][1:] == (2, 3, "")

    workbook = load_workbook(BytesIO(excel_bytes(result)))
    review = workbook["Review"]
    headers = [cell.value for cell in review[5]]
    for label, value in (
        ("Evidence Text", provider.candidates[0].evidence_text),
        ("Evidence URL", provider.candidates[0].evidence_url),
        ("Observed At", provider.candidates[0].observed_at),
        ("Status", "processing_failed"),
    ):
        assert review.cell(6, headers.index(label) + 1).value == value
    assert "Run: failed" in review["A2"].value
    assert workbook["People"]["A6"].value == "First"
    assert workbook["Companies"]["D8"].value == "pending"


def test_failed_batch_cannot_promote_partial_classification_to_people():
    provider = FailingDiscoveryProvider(candidates=[supported_candidate("Failure")])
    result = PeopleService(provider).run([CompanyTarget("Failure")], PeopleCriteria())

    assert not result.people
    assert len(result.review) == 1
    assert result.review[0]["Status"] == "processing_failed"
    assert "Discovery did not complete" in result.review[0]["Review Reason"]
    # The provider's input records remain unchanged for replay.
    assert provider.candidates[0].review_reason == ""


def test_failed_batch_does_not_promote_an_employment_unconfirmed_candidate():
    candidate = replace(supported_candidate("Failure"), employment_verified=False, employment_status="uncertain")
    provider = FailingDiscoveryProvider(candidates=[candidate])
    result = PeopleService(provider).run([CompanyTarget("Failure")], PeopleCriteria())
    assert not result.people and len(result.review) == 1
    assert result.review[0]["Status"] == "processing_failed"
    assert result.review[0]["Evidence Text"] == candidate.evidence_text
    assert "Discovery did not complete" in result.review[0]["Review Reason"]
