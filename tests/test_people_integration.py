from pathlib import Path
import json
import os
import subprocess
import sys
from threading import Event

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from company_gtm_enricher.people.excel_export import PEOPLE_COLUMNS
from company_gtm_enricher.people.job_runner import start_people_job
from company_gtm_enricher.people.models import CompanyTarget, PeopleCriteria
from company_gtm_enricher.people.providers import MockPeopleProvider


ROOT = Path(__file__).resolve().parents[1]


def cli(*args):
    return subprocess.run(
        [sys.executable, "-m", "company_gtm_enricher.cli", *args],
        cwd=ROOT, capture_output=True, text=True, timeout=15,
        env={**os.environ, "OPENAI_API_KEY": "", "SERPER_API_KEY": ""},
    )


def test_people_cli_exports_demo_and_company_cli_remains_csv(tmp_path):
    output = tmp_path / "demo.xlsx"
    run = cli("--mode", "people", "--provider", "mock", "--input", "examples/people_demo.csv", "--output", str(output))
    assert run.returncode == 0, run.stderr
    assert "synthetic" in run.stdout
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["People", "Companies", "Review"]
    assert workbook["People"].max_row == 9  # Four demo people after duplicate company merge.
    assert workbook["Companies"].max_row == 7
    assert "agent_execution" in workbook["Companies"].cell(6, 3).value
    assert workbook["Review"].max_row == 7
    csv_output = tmp_path / "companies.csv"
    old = cli("--provider", "mock", "--input", "examples/companies.csv", "--output", str(csv_output))
    assert old.returncode == 0, old.stderr
    assert "HQ City" in csv_output.read_text()


def test_people_cli_requires_keys_before_live_requests(tmp_path):
    output = tmp_path / "not_created.xlsx"
    run = cli("--mode", "people", "--provider", "openai", "--input", "examples/companies.csv", "--output", str(output))
    assert run.returncode == 1
    assert "OPENAI_API_KEY and SERPER_API_KEY" in run.stderr
    assert not output.exists()


def test_stopped_job_preserves_partial_and_pending_in_backup(tmp_path):
    entered, release = Event(), Event()

    class BlockingMock(MockPeopleProvider):
        def discover(self, company, criteria):
            entered.set()
            assert release.wait(5)
            return super().discover(company, criteria)

    job = start_people_job(
        [CompanyTarget("One"), CompanyTarget("Two")], PeopleCriteria(),
        BlockingMock(), PEOPLE_COLUMNS, "stopped.xlsx", tmp_path,
    )
    try:
        assert entered.wait(5)
        job.stop_event.set()
        release.set()
        job.thread.join(5)
        assert not job.thread.is_alive()
        assert job.status == "stopped"
        assert job.processed_companies == 1
        assert len(job.result.people) == 2
        assert job.result.companies[1]["Status"] == "pending"
        workbook = load_workbook(job.backup_path)
        assert "stopped" in workbook["Companies"]["A2"].value
        assert workbook["Companies"]["D7"].value == "pending"
    finally:
        job.stop_event.set()
        release.set()
        job.thread.join(5)


def test_provider_failure_is_visible_and_other_results_survive(tmp_path):
    class FailingMock(MockPeopleProvider):
        def discover(self, company, criteria):
            if company.name == "Broken":
                raise RuntimeError("demo failure")
            return super().discover(company, criteria)

    job = start_people_job(
        [CompanyTarget("Good"), CompanyTarget("Broken")], PeopleCriteria(),
        FailingMock(), PEOPLE_COLUMNS, "partial.xlsx", tmp_path,
    )
    job.thread.join(5)
    assert job.status == "completed_with_errors"
    assert len(job.result.people) == 2
    assert "demo failure" in job.result.companies[1]["Notes"]
    assert Path(job.backup_path).exists()


def test_people_ui_demo_and_return_to_company_mode(tmp_path, monkeypatch):
    # Keep demo backups inside this test's temporary directory.
    monkeypatch.chdir(tmp_path)
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=10).run()
    assert not app.exception
    app.sidebar.radio[0].set_value("People").run()
    assert not app.exception
    next(button for button in app.button if button.label == "Use demo companies").click().run()
    next(button for button in app.button if button.label == "Run people demo").click().run()
    job = app.session_state["people_job"]
    job.thread.join(5)
    app.run()
    assert not app.exception
    assert job.status == "completed"
    assert len(job.result.people) == 4
    assert len(job.result.review) == 2
    assert all(not person["LinkedIn URL"] for person in job.result.people)
    assert "synthetic" in app.info[0].value.lower()
    assert {tab.label for tab in app.tabs} == {"People", "Companies", "Review"}
    app.sidebar.radio[0].set_value("Company facts").run()
    assert not app.exception
    assert any(box.label == "Provider" for box in app.selectbox)


def test_custom_multiselect_criteria_are_preserved_and_live_requires_keys(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("OPENAI_API_KEY", "")
    monkeypatch.setenv("SERPER_API_KEY", "")
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=10).run()
    app.sidebar.radio[0].set_value("People").run()
    next(button for button in app.button if button.label == "Use demo companies").click().run()
    next(box for box in app.multiselect if box.label == "Departments").set_value(["Robotics"]).run()
    next(box for box in app.multiselect if box.label == "Leadership levels").set_value(["Research Lead"]).run()
    next(button for button in app.button if button.label == "Run people demo").click().run()
    job = app.session_state["people_job"]
    job.thread.join(5)
    app.run()
    assert not app.exception
    assert job.criteria.departments == ("Robotics",)
    assert job.criteria.leadership_levels == ("Research Lead",)
    assert len(job.result.people) == 6
    next(box for box in app.selectbox if box.label == "People source").set_value("Search + LLM").run()
    assert not app.exception
    assert next(button for button in app.button if button.label == "Find people").disabled
    assert any("SERPER_API_KEY" in warning.value for warning in app.warning)
    monkeypatch.setenv("OPENAI_API_KEY", "configured-for-test")
    monkeypatch.setenv("SERPER_API_KEY", "configured-for-test")
    app.run()
    assert next(button for button in app.button if button.label == "Find people").disabled
    assert any("fictional" in info.value for info in app.info)


def test_cli_json_keeps_settings_for_reproducible_evaluation(tmp_path):
    json_path = tmp_path / "run.json"
    run = cli(
        "--mode", "people", "--provider", "mock", "--input", "examples/people_demo.csv",
        "--output", str(tmp_path / "run.xlsx"), "--result-json", str(json_path),
        "--departments", "Infrastructure", "--leadership-levels", "Director",
    )
    assert run.returncode == 0, run.stderr
    result = json.loads(json_path.read_text())
    assert result["run_settings"]["criteria"]["leadership_levels"] == ["Director"]
    assert result["is_mock"]
    assert len(result["people"]) == 6
