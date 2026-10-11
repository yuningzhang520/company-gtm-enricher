from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from company_gtm_enricher.people.excel_export import PEOPLE_COLUMNS, excel_bytes
from company_gtm_enricher.people.models import PeopleResult


def read_result(**kwargs):
    return load_workbook(BytesIO(excel_bytes(PeopleResult(**kwargs))))


def value_cell(sheet, label, row=6):
    headers = [cell.value for cell in sheet[5]]
    return sheet.cell(row, headers.index(label) + 1)


def test_workbook_preserves_results_review_and_run_metadata():
    book = read_result(
        people=[{"Company": "Example", "Name": "Alex", "Status": "verified"}],
        companies=[{
            "Company": "Example", "Status": "partial", "People": 1, "Review": 1,
            "Notes": "Stopped after one record",
        }],
        review=[{"Name": "Sam", "Review Reason": "Employment needs confirmation"}],
        status="stopped", elapsed_seconds=1.234, provider_calls=2,
    )
    assert book.sheetnames == ["People", "Companies", "Review"]
    assert [cell.value for cell in book["People"][5]] == list(PEOPLE_COLUMNS)
    assert book["People"]["C6"].value == "Alex"
    assert book["Companies"]["E6"].value == 1
    assert book["Companies"]["D6"].value == "partial"
    assert book["Companies"]["G6"].value == "Stopped after one record"
    assert value_cell(book["Review"], "Review Reason").value == "Employment needs confirmation"
    for sheet in book:
        assert "stopped" in sheet["A2"].value
        assert "Provider calls: 2" in sheet["A2"].value
        assert "Elapsed: 1.23s" in sheet["A2"].value
        assert "not measured" in sheet["A3"].value
        assert sheet.freeze_panes == "A6"
        assert sheet.auto_filter.ref.endswith("6")


def test_empty_workbook_keeps_headers_and_synthetic_warning():
    book = read_result(people=[], companies=[], review=[], is_mock=True)
    for sheet in book:
        assert "SYNTHETIC DEMO DATA" in sheet["A1"].value
        assert "Not real people" in sheet["A1"].value
        assert "$0" in sheet["A3"].value
        assert "no paid API calls" in sheet["A3"].value
        assert sheet.max_row == 5
        assert sheet.auto_filter.ref.endswith("5")


def test_selected_people_columns_retain_status_and_review_diagnostics():
    result = PeopleResult(
        people=[{"Name": "Synthetic person", "Status": "mock"}], companies=[],
        review=[{"Review Reason": "Synthetic example"}], is_mock=True,
    )
    original = deepcopy(result)
    columns = ["Name"]
    book = load_workbook(BytesIO(excel_bytes(result, columns=columns)))
    assert result == original
    assert columns == ["Name"]
    assert [cell.value for cell in book["People"][5]] == ["Name", "Status"]
    assert book["People"]["B6"].value == "mock"
    assert "SYNTHETIC" in book["People"]["A1"].value
    assert [cell.value for cell in book["Review"][5]] == [*PEOPLE_COLUMNS, "Review Reason"]


@pytest.mark.parametrize("columns", [[], ["Name", "Name"], ["Unknown"]])
def test_invalid_columns_raise(columns):
    with pytest.raises(ValueError):
        excel_bytes(PeopleResult(people=[], companies=[], review=[]), columns)


@pytest.mark.parametrize("url", [
    "https://www.linkedin.com/in/example", "http://example.com/evidence",
    "javascript:alert(1)", "file:///tmp/private", "https://", "https://[invalid",
    "https://user:secret@example.com", "https://example.com:invalid", "https://example.com/ bad",
])
def test_only_safe_web_urls_are_clickable(url):
    book = read_result(people=[{"LinkedIn URL": url, "Evidence URL": url}], companies=[], review=[])
    safe = url in {"https://www.linkedin.com/in/example", "http://example.com/evidence"}
    for address in ("E6", "I6"):
        cell = book["People"][address]
        assert cell.value == url
        assert bool(cell.hyperlink) is safe
        if safe:
            assert cell.hyperlink.target == url


@pytest.mark.parametrize("value", ['=HYPERLINK("https://example.com","click")', "+cmd", "-1+2", "@SUM(1)"])
def test_user_values_are_literal_text_and_xml_controls_are_removed(value):
    book = read_result(
        people=[{"Name": value, "Notes": "before\x00\x0bafter\nnext"}],
        companies=[{"Company": value}], review=[{"Review Reason": value}],
    )
    for sheet, label in (("People", "Name"), ("Companies", "Company"), ("Review", "Review Reason")):
        assert value_cell(book[sheet], label).value == value
        assert value_cell(book[sheet], label).data_type == "s"
    assert book["People"]["J6"].value == "beforeafter\nnext"
