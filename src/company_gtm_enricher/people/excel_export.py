"""Export people results, review evidence, and run metadata as a safe workbook."""
from __future__ import annotations

from io import BytesIO
import json
from typing import Sequence
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import PeopleResult

PEOPLE_COLUMNS = (
    "Company", "List Tags", "Name", "Title", "LinkedIn URL", "Department",
    "Seniority", "Status", "Evidence URL", "Notes", "Domain", "Evidence Text", "Observed At",
)
COMPANY_COLUMNS = ("Company", "Domain", "List Tags", "Status", "People", "Review", "Notes")
REVIEW_COLUMNS = (*PEOPLE_COLUMNS, "Review Reason")
HEADER_ROW = 5


def _safe_url(value: str) -> bool:
    if any(char.isspace() or ord(char) < 32 for char in value) or "\\" in value:
        return False
    try:
        url = urlsplit(value)
        return (
            url.scheme.lower() in {"https", "http"}
            and bool(url.hostname)
            and url.username is None
            and url.password is None
            and (url.port is None or 0 < url.port < 65536)
        )
    except ValueError:
        return False


def _write(cell, value) -> None:
    if value is None:
        return
    if isinstance(value, (int, float, bool)):
        cell.value = value
    else:
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(value))
        # Explicit string type preserves formula-like input without evaluating it.
        cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def excel_bytes(result: PeopleResult, columns: Sequence[str] | None = None) -> bytes:
    selected = list(PEOPLE_COLUMNS if columns is None else columns)
    if not selected or len(set(selected)) != len(selected):
        raise ValueError("Choose at least one output column, without duplicates.")
    if any(column not in PEOPLE_COLUMNS for column in selected):
        raise ValueError("Unknown people output column.")
    if "Status" not in selected:
        selected.append("Status")

    workbook = Workbook()
    workbook.remove(workbook.active)
    banner = (
        "SYNTHETIC DEMO DATA — Not real people or verified research"
        if result.is_mock else "People enrichment — inspect status and supporting evidence"
    )
    metadata = (
        banner,
        f"Run: {result.status} | Provider calls: {result.provider_calls} | "
        f"Elapsed: {result.elapsed_seconds:.2f}s",
        "API charges: $0 — offline mock; no paid API calls."
        if result.is_mock else (
            f"API charges: not measured. Search requests: {result.usage.get('search_requests', 0)}; "
            f"LLM requests: {result.usage.get('llm_requests', 0)}; "
            f"Company page requests: {result.usage.get('page_requests', 0)}; "
            f"input/output tokens: {result.usage.get('input_tokens', 0)}/{result.usage.get('output_tokens', 0)}"
        ),
        json.dumps(result.run_settings, ensure_ascii=False) if result.run_settings else "",
    )
    widths = {"Title": 38, "Notes": 48, "Review Reason": 48, "LinkedIn URL": 48, "Evidence URL": 48, "Evidence Text": 65, "Observed At": 28}
    for name, headers, records in (
        ("People", selected, result.people),
        ("Companies", COMPANY_COLUMNS, result.companies),
        ("Review", REVIEW_COLUMNS, result.review),
    ):
        sheet = workbook.create_sheet(name)
        for row, message in enumerate(metadata, 1):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
            _write(sheet.cell(row, 1), message)
            sheet.cell(row, 1).font = Font(bold=row == 1)
            sheet.row_dimensions[row].height = 32
        sheet.cell(1, 1).fill = PatternFill("solid", fgColor="FFF2CC" if result.is_mock else "DDEBF7")
        for index, header in enumerate(headers, 1):
            cell = sheet.cell(HEADER_ROW, index)
            _write(cell, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244062")
            sheet.column_dimensions[get_column_letter(index)].width = widths.get(header, 24)
        for row, record in enumerate(records, HEADER_ROW + 1):
            for index, header in enumerate(headers, 1):
                cell = sheet.cell(row, index)
                value = record.get(header, "")
                _write(cell, value)
                if header in {"LinkedIn URL", "Evidence URL"} and isinstance(value, str) and _safe_url(value):
                    cell.hyperlink = value
                    cell.font = Font(color="0563C1", underline="single")
        sheet.freeze_panes = f"A{HEADER_ROW + 1}"
        sheet.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(headers))}{max(HEADER_ROW, sheet.max_row)}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
